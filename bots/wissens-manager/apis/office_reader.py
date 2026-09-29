"""
Office-Reader – liest .docx und .xlsx
"""

import os


class OfficeReader:
    def extract_docx(self, path):
        """Liest Text aus .docx"""
        try:
            from docx import Document
            doc = Document(path)
            
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            return "\n".join(text_parts)
        except Exception as e:
            print(f"   ❌ DOCX-Fehler: {e}")
            return None
    
    def extract_xlsx(self, path):
        """Liest Text aus .xlsx"""
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            
            text_parts = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_parts.append(f"=== {sheet_name} ===")
                
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(c) for c in row if c is not None])
                    if row_text.strip():
                        text_parts.append(row_text)
            
            wb.close()
            return "\n".join(text_parts)
        except Exception as e:
            print(f"   ❌ XLSX-Fehler: {e}")
            return None
    
    def extract_text(self, path):
        """Erkennt Typ und ruft passende Methode auf"""
        ext = os.path.splitext(path)[1].lower()
        
        if ext == ".docx":
            return self.extract_docx(path)
        elif ext == ".xlsx":
            return self.extract_xlsx(path)
        elif ext in [".txt", ".md"]:
            try:
                with open(path, "r", encoding="utf-8") as f:
                    return f.read()
            except UnicodeDecodeError:
                with open(path, "r", encoding="latin-1") as f:
                    return f.read()
        
        return None